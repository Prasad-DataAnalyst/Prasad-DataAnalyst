"""
Builds the XLSX test workbooks of the formats-xlsx module (src/formats/xlsx/__fixtures__).

Each workbook is written with openpyxl and then re-saved by LibreOffice headless, so formulas carry
cached values and the files come from a second, independent producer, as SAD section 13.2 asks for
the messy-file library. The committed .xlsx files are the test inputs; this script documents how
they were made and can rebuild them:

    python3 src/formats/xlsx/__fixtures__/make_fixtures.py

LibreOffice stamps its own save time into docProps, so a rebuild changes the bytes but not the
content the tests read. Author: Prasad (Prasad-DataAnalyst).
"""

import datetime as dt
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

HERE = os.path.dirname(os.path.abspath(__file__))


def hidden_sheets_formulas(path):
    wb = Workbook()
    notes = wb.active
    notes.title = "Calc notes"
    notes["A1"] = "internal"
    notes["B1"] = 1
    notes.sheet_state = "hidden"

    prod = wb.create_sheet("Production")
    prod.append(["Well", "Oil (bbl)", "Water (bbl)", "Water cut", "Oil per water", "Label", "Big producer"])
    rows = [("A-01", 100, 25), ("A-02", 80, 20), ("A-03", 50, 0), ("A-04", 0, 0)]
    for i, (well, oil, water) in enumerate(rows, start=2):
        prod.append([
            well,
            oil,
            water,
            f"=C{i}/(B{i}+C{i})",
            f"=B{i}/C{i}",
            f'=A{i}&"-ok"',
            f"=B{i}>60",
        ])

    lookup = wb.create_sheet("Lookup")
    lookup["A1"] = "code"
    lookup["B1"] = "value"
    lookup.sheet_state = "veryHidden"

    summary = wb.create_sheet("Summary")
    summary["A1"] = "Total oil"
    summary["B1"] = "=SUM(Production!B2:B5)"
    wb.active = 1
    wb.save(path)


def date1904(path):
    wb = Workbook()
    wb.epoch = CALENDAR_MAC_1904
    ws = wb.active
    ws.title = "Readings"
    ws.append(["Date", "Value"])
    ws.append([dt.date(2026, 1, 31), 10.5])
    ws.append([dt.date(2026, 2, 1), 11.25])
    ws.append([dt.datetime(2026, 2, 2, 6, 30), 12])
    for row in range(2, 4):
        ws.cell(row=row, column=1).number_format = "yyyy-mm-dd"
    ws.cell(row=4, column=1).number_format = "yyyy-mm-dd hh:mm"
    wb.save(path)


def text_and_serial_dates(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily"
    ws.append(["Date", "Date text", "Oil", "Timestamp"])
    data = [
        (dt.date(2026, 1, 31), "31/01/2026", 101.5, dt.datetime(2026, 1, 31, 6, 0, 0)),
        (dt.date(2026, 2, 1), "2026-02-01", 99.0, dt.datetime(2026, 2, 1, 6, 0, 30)),
        (dt.date(2026, 2, 2), "02-Feb-2026", 98.25, dt.datetime(2026, 2, 2, 18, 45, 0)),
    ]
    for row in data:
        ws.append(list(row))
    for row in range(2, 5):
        ws.cell(row=row, column=1).number_format = "dd/mm/yyyy"
        ws.cell(row=row, column=4).number_format = "yyyy-mm-dd hh:mm:ss"
    wb.save(path)


def numbers_as_text(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "IDs"
    ws.append(["Well ID", "Rate text", "Rate"])
    ws.append(["00123", "12.5", 12.5])
    ws.append(["00456", "1,234.5", 1234.5])
    ws.append(["07", "7", 7])
    wb.save(path)


def merged_two_row_header_units(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Well tests"
    ws["A1"] = "Well test summary - Field Alpha"
    ws.merge_cells("A1:F1")
    ws["A2"] = "Well"
    ws.merge_cells("A2:A3")
    ws["B2"] = "Test date"
    ws.merge_cells("B2:B3")
    ws["C2"] = "Pressure"
    ws.merge_cells("C2:D2")
    ws["E2"] = "Temperature"
    ws.merge_cells("E2:F2")
    for col, text in zip("CDEF", ["Max", "Min", "Max", "Min"]):
        ws[f"{col}3"] = text
    for col, text in zip("CDEF", ["psi", "psi", "degF", "degF"]):
        ws[f"{col}4"] = text
    data = [
        ("A-01", dt.date(2026, 1, 5), 2510.5, 2380.0, 182.5, 170.0),
        ("A-02", dt.date(2026, 1, 6), 2475.0, 2301.25, 179.0, 168.5),
        ("A-03", dt.date(2026, 1, 7), 2602.0, 2444.0, 185.25, 174.0),
        ("B-01", dt.date(2026, 1, 8), 1990.0, 1850.5, 160.0, 151.0),
        ("B-02", dt.date(2026, 1, 9), 2050.75, 1902.0, 163.5, 155.25),
    ]
    for r, row in enumerate(data, start=5):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if c == 2:
                cell.number_format = "yyyy-mm-dd"
    wb.save(path)


def two_tables_one_sheet(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["A1"] = "Daily production report"
    ws.append([])
    ws.append(["Well", "Date", "Oil [bbl/d]", "Gas [Mscf/d]"])
    production = [
        ("A-01", dt.date(2026, 3, 1), 1200.5, 850.0),
        ("A-02", dt.date(2026, 3, 1), 980.0, 701.25),
        ("A-03", dt.date(2026, 3, 1), 1105.75, 790.5),
        ("B-01", dt.date(2026, 3, 1), 640.0, 455.0),
        ("B-02", dt.date(2026, 3, 1), 702.25, 498.75),
    ]
    for r, row in enumerate(production, start=4):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if c == 2:
                cell.number_format = "yyyy-mm-dd"
    ws["A11"] = "Well tests"
    for c, text in enumerate(["Well", "Test date", "Rate [bbl/d]"], start=1):
        ws.cell(row=12, column=c, value=text)
    tests = [("A-01", dt.date(2026, 2, 20), 1250.0), ("A-02", dt.date(2026, 2, 21), 990.5), ("B-01", dt.date(2026, 2, 22), 655.0)]
    for r, row in enumerate(tests, start=13):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if c == 2:
                cell.number_format = "yyyy-mm-dd"
    wb.save(path)


BUILDERS = {
    "xlsx_hidden_sheets_formulas.xlsx": hidden_sheets_formulas,
    "xlsx_date1904.xlsx": date1904,
    "xlsx_text_and_serial_dates.xlsx": text_and_serial_dates,
    "xlsx_numbers_as_text.xlsx": numbers_as_text,
    "xlsx_merged_two_row_header_units.xlsx": merged_two_row_header_units,
    "xlsx_two_tables_one_sheet.xlsx": two_tables_one_sheet,
}


def resave_with_libreoffice(src, out_dir):
    """LibreOffice computes the formula results that openpyxl leaves out and writes the file again."""
    profile = tempfile.mkdtemp(prefix="tw-lo-profile-")
    try:
        subprocess.run(
            [
                "soffice",
                f"-env:UserInstallation=file://{profile}",
                "--headless",
                "--convert-to",
                "xlsx:Calc MS Excel 2007 XML",
                "--outdir",
                out_dir,
                src,
            ],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=120,
        )
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def main():
    work = tempfile.mkdtemp(prefix="tw-xlsx-fixtures-")
    try:
        raw_dir = os.path.join(work, "raw")
        out_dir = os.path.join(work, "out")
        os.makedirs(raw_dir)
        os.makedirs(out_dir)
        for name, build in BUILDERS.items():
            raw = os.path.join(raw_dir, name)
            build(raw)
            resave_with_libreoffice(raw, out_dir)
            shutil.copyfile(os.path.join(out_dir, name), os.path.join(HERE, name))
            print(f"wrote {name}")
    finally:
        shutil.rmtree(work, ignore_errors=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
