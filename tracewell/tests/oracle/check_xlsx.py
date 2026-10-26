"""Excel oracle: python3 -I check_xlsx.py <package.zip>. Opens data.xlsx with openpyxl and reports
sheet names, frozen panes, filters, merged cells, cell types and the SAD 9.4 digest of every column."""

import datetime
import io
import json
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from canonical import column_digest  # noqa: E402
from common import emit, open_package  # noqa: E402


def ms_iso(value):
    """A datetime to the millisecond, as the engine stores datetimes (Excel keeps no zone: UTC)."""
    micro = value.microsecond
    rounded = value.replace(microsecond=0) + datetime.timedelta(milliseconds=round(micro / 1000))
    return rounded.strftime("%Y-%m-%dT%H:%M:%S.") + "%03d" % (rounded.microsecond // 1000) + "Z"


def cell_value(cell, kind):
    v = cell.value
    if v is None or v == "":
        return None
    if kind == "date":
        if isinstance(v, datetime.datetime):
            return {"date": v.date().isoformat()}
        if isinstance(v, datetime.date):
            return {"date": v.isoformat()}
        raise ValueError(f"{cell.coordinate} is not a date: {v!r}")
    if kind == "datetime":
        if isinstance(v, datetime.datetime):
            return {"datetime": ms_iso(v)}
        raise ValueError(f"{cell.coordinate} is not a datetime: {v!r}")
    if kind in ("number", "integer"):
        if isinstance(v, bool) or not isinstance(v, (int, float)):
            raise ValueError(f"{cell.coordinate} is not a number: {v!r}")
        return float(v)
    if kind == "boolean":
        if not isinstance(v, bool):
            raise ValueError(f"{cell.coordinate} is not a boolean: {v!r}")
        return v
    if kind == "geometry":
        return {"geometry": json.loads(v)}
    return str(v)


def main(path):
    files, manifest = open_package(path)
    book = openpyxl.load_workbook(io.BytesIO(files["data.xlsx"]))
    sheet = book["Data"]
    columns = manifest["dataset"]["columns"]
    rows = list(sheet.iter_rows(min_row=1))
    header = [c.value for c in rows[0]] if rows else []
    body = rows[1:]
    digests, types = [], []
    for i, col in enumerate(columns):
        values = [cell_value(r[i], col["type"]) for r in body]
        digests.append({"header": header[i] if i < len(header) else None, "sha256": column_digest(values)})
        sample = next((r[i] for r in body if r[i].value is not None), None)
        types.append(
            {
                "header": header[i] if i < len(header) else None,
                "data_type": None if sample is None else sample.data_type,
                "is_date": None if sample is None else bool(sample.is_date),
            }
        )
    emit(
        {
            "sheets": book.sheetnames,
            "freeze_panes": sheet.freeze_panes,
            "auto_filter": sheet.auto_filter.ref,
            "merged": [str(m) for m in sheet.merged_cells.ranges],
            "rows": len(body),
            "header": header,
            "types": types,
            "columns": digests,
            "epoch1904": bool(book.epoch == openpyxl.utils.datetime.CALENDAR_MAC_1904),
        }
    )


if __name__ == "__main__":
    main(sys.argv[1])
